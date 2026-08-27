import csv
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path(__file__).with_name("tt.txt")
OUTPUT = SOURCE.with_suffix(".xlsx")


def main():
    text = SOURCE.read_text(encoding="utf-8-sig")
    chunks = re.split(r" (?=\d{9},)", text)
    header = next(csv.reader([chunks[0]]))
    rows = [next(csv.reader([chunk])) for chunk in chunks[1:]]

    if not rows:
        raise ValueError("No data rows found")
    invalid = [index for index, row in enumerate(rows, start=2) if len(row) != len(header)]
    if invalid:
        raise ValueError(f"Rows with unexpected column counts: {invalid}")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "보험 데이터"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    sheet.sheet_view.zoomScale = 85

    sheet.append(header)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    # Preserve identifiers, resident numbers, dates, and leading zeroes exactly.
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="center")

    for column_index, column_name in enumerate(header, start=1):
        values = [column_name] + [row[column_index - 1] for row in rows]
        width = min(max(len(str(value)) for value in values) + 2, 28)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 10)

    workbook.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=True, data_only=True)
    check_sheet = check.active
    if check_sheet.max_row != len(rows) + 1 or check_sheet.max_column != len(header):
        raise RuntimeError("Excel validation failed")
    check.close()
    print(f"Created {OUTPUT.name}: {len(rows)} rows, {len(header)} columns")


if __name__ == "__main__":
    main()
